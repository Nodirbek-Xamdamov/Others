"""Excel (Google Sheets eksporti) -> in-memory SQLite.

Har bir varaq tozalanadi va aniq nomli ustunlarga ega jadvalga aylanadi.
Yuk mashinalari haqidagi 6 ta varaq bitta `all_trucks` jadvaliga yig'iladi
(`sheet` ustuni qaysi varaqdan kelganini ko'rsatadi), LLM esa undan olingan
view'larni ko'radi: `trucks` = faqat aktiv park (TRUCK INFO), `rental_trucks`,
`other_truck_lists`, `inactive_trucks`. Shunda "nechta mashina" savoliga
arxiv va takroriy ro'yxatlar qo'shilib ketmaydi.

Tanish bo'lmagan yangi varaq qo'shilsa, u avtomatik ravishda (1-qator = sarlavha)
alohida jadval sifatida yuklanadi.
"""
from __future__ import annotations

import datetime as dt
import re
import sqlite3
import warnings
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

# --------------------------------------------------------------------------- #
# Qiymatlarni tozalash
# --------------------------------------------------------------------------- #

_WS = re.compile(r"\s+")
_VIN = re.compile(r"^[A-Z0-9]{16,18}$")
# "Driver" ustunida haydovchi o'rniga joy/holat yozilgan qiymatlar
_PLACEHOLDER = re.compile(
    r"(?i)^(at\b|alt shop|new\s+yard|.*\byard\b|.*\bshop\b|inactive|returned|"
    r"needs new driver|no driver|home time|oos\b|out of service|in?cative|.*\bdealer\b|.*\bmtp\b|.*\bttc\b)"
)
_EMPTY = {"", "-", "--", "n/a", "na", "none", "null", "#n/a", "#value!", "#ref!", "#div/0!"}


def clean_text(v) -> str | None:
    """Istalgan qiymatni toza matnga aylantiradi (ID'lar uchun: 417.0 -> '417')."""
    if v is None:
        return None
    if isinstance(v, bool):
        return "yes" if v else "no"
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, (dt.datetime, dt.date)):
        return to_date(v)
    s = _WS.sub(" ", str(v)).strip()
    s = s.replace("☑️", "").replace("☑", "").strip()
    return None if s.lower() in _EMPTY else s


def to_int(v) -> int | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(round(v))
    s = re.sub(r"[,\s$]", "", str(v))
    try:
        return int(round(float(s)))
    except ValueError:
        return None


def to_float(v) -> float | None:
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(re.sub(r"[,\s$%]", "", str(v)))
    except ValueError:
        return None


def to_date(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0) else v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d", "%m-%d-%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def norm_name(v) -> str | None:
    s = clean_text(v)
    if s is None:
        return None
    return s.title() if (s.isupper() or s.islower()) and len(s) > 3 else s


def split_driver(v) -> tuple[str | None, str | None]:
    """'At MTP' kabi qiymatlarni haydovchi emas, joylashuv/holat deb ajratadi."""
    s = clean_text(v)
    if s is None:
        return None, None
    if _PLACEHOLDER.match(s):
        return None, s
    return s, None


def is_vin(v) -> bool:
    s = clean_text(v)
    return bool(s and _VIN.match(s.upper().replace(" ", "")) and re.search(r"\d", s) and re.search(r"[A-Z]", s.upper()))


def cell(row, i):
    return row[i] if 0 <= i < len(row) else None


def nonempty(row) -> bool:
    return any(clean_text(v) is not None for v in row)


# --------------------------------------------------------------------------- #
# Jadval ta'riflari
# --------------------------------------------------------------------------- #

@dataclass
class Table:
    name: str
    description: str
    columns: list[tuple[str, str, str]]  # (nom, sqlite tipi, tavsif)
    rows: list[dict] = field(default_factory=list)


TRUCK_COLUMNS = [
    ("sheet", "TEXT", "qaysi varaqdan: 'TRUCK INFO' (asosiy aktiv park), 'CAC Trucks', 'DEMIR Trucks', 'New trucks', 'Rental trucks', 'Inactive Trucks'"),
    ("driver", "TEXT", "haydovchi ismi (Inactive'da oxirgi haydovchi). NULL bo'lsa haydovchi yo'q"),
    ("coworker", "TEXT", "ikkinchi haydovchi (team)"),
    ("location_note", "TEXT", "haydovchi o'rniga yozilgan joy/holat: 'At MTP', 'NEW YARD', 'AT SHOP', 'At TTC YARD', 'Returned', 'Needs new driver' ..."),
    ("unit_number", "TEXT", "yuk mashinasi raqami (unit), masalan '417', '9693 JCE', '567282 RYL'"),
    ("vehicle_group", "TEXT", "guruh/kompaniya (Riverview, Jet Cargo Express, Ryder Lease ...)"),
    ("owner", "TEXT", "egasi / provider"),
    ("color", "TEXT", "rangi"),
    ("make", "TEXT", "ishlab chiqaruvchi (Freightliner, Volvo, International ...)"),
    ("model", "TEXT", "model (Cascadia, VNL, LT 625 ...)"),
    ("year", "INTEGER", "yili"),
    ("vin", "TEXT", "VIN raqami"),
    ("plate_number", "TEXT", "davlat raqami (plate)"),
    ("prepass", "TEXT", "PrePass transponder ID"),
    ("ifta", "TEXT", "IFTA stiker raqami"),
    ("cab_card_state", "TEXT", "cab card (registratsiya) shtati"),
    ("cab_card_expiration", "DATE", "cab card (registratsiya) tugash sanasi"),
    ("inspection_state", "TEXT", "yillik inspeksiya shtati"),
    ("inspection_expiration", "DATE", "yillik inspeksiya tugash sanasi"),
    ("latest_pm_miles", "INTEGER", "oxirgi PM (texnik xizmat, moy almashtirish) qilingan probeg"),
    ("next_pm_due_miles", "INTEGER", "keyingi PM qilinishi kerak bo'lgan probeg"),
    ("miles_until_pm", "INTEGER", "keyingi PM gacha qolgan mil (next_pm_due_miles - current_mileage). MANFIY = PM muddati o'tib ketgan"),
    ("current_mileage", "INTEGER", "joriy probeg (mil)"),
    ("status_note", "TEXT", "qo'shimcha izoh (masalan 'ALL GOOD')"),
]


def _truck_row(sheet: str, r, *, driver_i, cow_i, unit_i, group_i, owner_i, anchor_vin=9) -> dict:
    """TRUCK INFO uslubidagi qator. Ustunlar VIN joyiga nisbatan olinadi,
    chunki ba'zi qatorlarda ustunlar 1 taga surilgan."""
    vin_i = next((i for i in range(anchor_vin - 1, anchor_vin + 3) if is_vin(cell(r, i))), anchor_vin)
    o = vin_i - anchor_vin
    driver, note = split_driver(cell(r, driver_i))
    return {
        "sheet": sheet,
        "driver": norm_name(driver),
        "coworker": norm_name(cell(r, cow_i)) if cow_i is not None else None,
        "location_note": note,
        "unit_number": clean_text(cell(r, unit_i)),
        "vehicle_group": clean_text(cell(r, group_i)) if group_i is not None else None,
        "owner": clean_text(cell(r, owner_i)) if owner_i is not None else None,
        "color": norm_name(cell(r, 5 + o)),
        "make": norm_name(cell(r, 6 + o)),
        "model": clean_text(cell(r, 7 + o)),
        "year": to_int(cell(r, 8 + o)),
        "vin": clean_text(cell(r, 9 + o)),
        "plate_number": clean_text(cell(r, 10 + o)),
        "prepass": clean_text(cell(r, 11 + o)),
        "ifta": clean_text(cell(r, 12 + o)),
        "cab_card_state": clean_text(cell(r, 13 + o)),
        "cab_card_expiration": to_date(cell(r, 14 + o)),
        "inspection_state": clean_text(cell(r, 15 + o)),
        "inspection_expiration": to_date(cell(r, 16 + o)),
        "latest_pm_miles": to_int(cell(r, 17 + o)),
        "next_pm_due_miles": to_int(cell(r, 18 + o)),
        "miles_until_pm": to_int(cell(r, 19 + o)),
        "current_mileage": to_int(cell(r, 20 + o)),
    }


def parse_standard_trucks(ws, sheet: str) -> list[dict]:
    """TRUCK INFO, CAC Trucks, DEMIR Trucks."""
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r) or clean_text(cell(r, 2)) is None:
            continue
        row = _truck_row(sheet, r, driver_i=0, cow_i=1, unit_i=2, group_i=3, owner_i=4)
        row["status_note"] = clean_text(cell(r, 23))
        out.append(row)
    return out


def parse_inactive(ws, sheet: str) -> list[dict]:
    """Inactive Trucks: arxiv, qatorlar ikki xil formatda."""
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r):
            continue
        first = cell(r, 0)
        if isinstance(first, str) and clean_text(first):  # TRUCK INFO formati
            row = _truck_row(sheet, r, driver_i=0, cow_i=1, unit_i=2, group_i=3, owner_i=4)
        else:  # eski format: #, Last driver, coworker, unit, provider, color...
            row = _truck_row(sheet, r, driver_i=1, cow_i=2, unit_i=3, group_i=None, owner_i=4)
        if row["unit_number"] is None:
            continue
        row["status_note"] = clean_text(cell(r, 23))
        out.append(row)
    return out


def parse_new_trucks(ws, sheet: str) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r) or clean_text(cell(r, 2)) is None:
            continue
        driver, note = split_driver(cell(r, 0))
        out.append({
            "sheet": sheet, "driver": norm_name(driver), "coworker": norm_name(cell(r, 1)),
            "location_note": note, "unit_number": clean_text(cell(r, 2)),
            "vehicle_group": clean_text(cell(r, 3)), "color": norm_name(cell(r, 4)),
            "make": norm_name(cell(r, 5)), "model": clean_text(cell(r, 6)), "year": to_int(cell(r, 7)),
            "vin": clean_text(cell(r, 8)), "plate_number": clean_text(cell(r, 9)),
            "latest_pm_miles": to_int(cell(r, 10)), "next_pm_due_miles": to_int(cell(r, 11)),
            "miles_until_pm": to_int(cell(r, 12)), "current_mileage": to_int(cell(r, 13)),
        })
    return out


def parse_rental(ws, sheet: str) -> list[dict]:
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r) or clean_text(cell(r, 2)) is None:
            continue
        driver, note = split_driver(cell(r, 0))
        out.append({
            "sheet": sheet, "driver": norm_name(driver), "coworker": norm_name(cell(r, 1)),
            "location_note": note, "unit_number": clean_text(cell(r, 2)),
            "vehicle_group": clean_text(cell(r, 3)), "make": norm_name(cell(r, 4)),
            "model": clean_text(cell(r, 5)), "year": to_int(cell(r, 6)),
            "vin": clean_text(cell(r, 7)),  # varaqda sarlavha 'Mileage' deb xato yozilgan, aslida VIN
        })
    return out


def parse_recent_issues(ws) -> Table:
    t = Table("recent_issues", "Varaq 'Sheet103': mashinalardagi oxirgi nosozlik/ta'mirlar", [
        ("driver", "TEXT", "haydovchi"), ("coworker", "TEXT", "ikkinchi haydovchi"),
        ("location_note", "TEXT", "haydovchi o'rniga joy ('At TTC', 'At MTP')"),
        ("unit_number", "TEXT", "unit raqami"), ("make", "TEXT", ""), ("model", "TEXT", ""),
        ("year", "INTEGER", ""), ("vin", "TEXT", ""), ("mileage", "INTEGER", "probeg"),
        ("recent_issue", "TEXT", "oxirgi muammo/ta'mir, masalan 'Radiator replacement'"),
    ])
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r) or clean_text(cell(r, 2)) is None:
            continue
        driver, note = split_driver(cell(r, 0))
        t.rows.append({
            "driver": norm_name(driver), "coworker": norm_name(cell(r, 1)), "location_note": note,
            "unit_number": clean_text(cell(r, 2)), "make": norm_name(cell(r, 3)), "model": clean_text(cell(r, 4)),
            "year": to_int(cell(r, 5)), "vin": clean_text(cell(r, 6)), "mileage": to_int(cell(r, 7)),
            "recent_issue": clean_text(cell(r, 8)),
        })
    return t


def parse_telematics(ws) -> Table:
    t = Table("gps_tracking", "Varaq 'Insurance VIN': Samsara GPS eksporti (mashina qayerda, yoqilg'i, holati)", [
        ("unit_number", "TEXT", "unit raqami"), ("gateway_serial", "TEXT", "GPS qurilma seriyasi"),
        ("location", "TEXT", "oxirgi joylashuv (shahar, shtat yoki 'TTC Yard')"),
        ("last_connected", "TEXT", "oxirgi ulanish vaqti"),
        ("last_trip", "TEXT", "oxirgi reys vaqti yoki 'On trip' (hozir yo'lda)"),
        ("engine_status", "TEXT", "'on', 'off', 'idle'"), ("fuel_level_pct", "REAL", "yoqilg'i darajasi %"),
        ("current_driver", "TEXT", ""), ("license_plate", "TEXT", ""),
        ("make", "TEXT", ""), ("model", "TEXT", ""), ("year", "INTEGER", ""), ("vin", "TEXT", ""),
        ("photos_url", "TEXT", "Samsara rasm havolasi"),
    ])
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not nonempty(r) or clean_text(cell(r, 0)) is None:
            continue
        lt = cell(r, 4)
        t.rows.append({
            "unit_number": clean_text(cell(r, 0)), "gateway_serial": clean_text(cell(r, 1)),
            "location": clean_text(cell(r, 2)), "last_connected": to_date(cell(r, 3)),
            "last_trip": to_date(lt) if isinstance(lt, dt.datetime) else clean_text(lt),
            "engine_status": clean_text(cell(r, 5)), "fuel_level_pct": to_float(cell(r, 6)),
            "current_driver": norm_name(cell(r, 7)), "license_plate": clean_text(cell(r, 8)),
            "make": norm_name(cell(r, 11)), "model": norm_name(cell(r, 12)), "year": to_int(cell(r, 13)),
            "vin": clean_text(cell(r, 14)), "photos_url": clean_text(cell(r, 15)),
        })
    return t


PTI_STATUSES = {"need ask", "no need ask", "asked / done", "asked / didn't send"}


def parse_pti(ws) -> Table:
    t = Table("pti_checks", "Varaq 'PTI': haydovchilardan kunlik PTI (pre-trip inspection) so'ralganmi. Har qator = bitta haydovchi/team + bitta kun", [
        ("drivers", "TEXT", "haydovchi yoki team ('Ism & Ism CD')"),
        ("note", "TEXT", "izoh, masalan 'OFF'"),
        ("check_date", "DATE", "sana"),
        ("status", "TEXT", "'Need ask' (so'rash kerak), 'No Need ASK' (kerak emas), 'Asked / Done' (so'raldi, yuborildi), 'Asked / Didn't send' (so'raldi, yubormadi)"),
    ])
    dates: list = []
    for r in ws.iter_rows(values_only=True):
        if clean_text(cell(r, 0)) == "Trucks":
            dates = [to_date(v) if isinstance(v, dt.datetime) else None for v in r[2:9]]
            continue
        name = clean_text(cell(r, 0))
        if not name or not dates:
            continue
        for d, v in zip(dates, r[2:9]):
            s = clean_text(v)
            if d and s and s.lower() in PTI_STATUSES:
                t.rows.append({"drivers": name, "note": clean_text(cell(r, 1)), "check_date": d, "status": s})
    return t


def parse_contacts(ws) -> Table:
    t = Table("contacts", "Varaq 'Sheet45': ofis xodimlari kontaktlari (sarlavhasiz)", [
        ("name", "TEXT", ""), ("email", "TEXT", ""), ("phone", "TEXT", ""),
    ])
    for r in ws.iter_rows(values_only=True):
        if nonempty(r) and clean_text(cell(r, 0)):
            t.rows.append({"name": norm_name(cell(r, 0)), "email": clean_text(cell(r, 1)), "phone": clean_text(cell(r, 2))})
    return t


def parse_deductions(ws) -> Table:
    t = Table("deductions_2022", "Varaq '2022': 2022 yilgi broker jarimalari/ushlab qolishlar (sarlavhasiz; ustun ma'nolari taxminiy)", [
        ("reason", "TEXT", "sabab"), ("date", "DATE", ""), ("load_number", "TEXT", "yuk raqami"),
        ("amount_deducted", "REAL", "ushlab qolingan summa $"), ("load_amount", "REAL", "yuk narxi $ (taxminan)"),
    ])
    for r in ws.iter_rows(values_only=True):
        if nonempty(r) and clean_text(cell(r, 0)):
            t.rows.append({"reason": clean_text(cell(r, 0)), "date": to_date(cell(r, 1)), "load_number": clean_text(cell(r, 2)),
                           "amount_deducted": to_float(cell(r, 4)), "load_amount": to_float(cell(r, 5))})
    return t


def parse_days_off(ws) -> Table:
    t = Table("staff_days_off", "Varaq 'Sheet35': xodimlarning dam olish kunlari", [
        ("name", "TEXT", ""), ("day_off", "TEXT", "hafta kuni (Monday..Sunday)"),
    ])
    header = [clean_text(v) for v in next(ws.iter_rows(max_row=1, values_only=True))]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not clean_text(cell(r, 0)):
            continue
        for day, v in zip(header[1:], r[1:]):
            if day and clean_text(v) and clean_text(v).upper() == "OFF":
                t.rows.append({"name": clean_text(cell(r, 0)), "day_off": day})
    return t


def parse_generic(ws) -> Table | None:
    """Noma'lum varaq: birinchi to'ldirilgan qator sarlavha."""
    rows = [r for r in ws.iter_rows(values_only=True) if nonempty(r)]
    if len(rows) < 2:
        return None
    header = rows[0]
    cols, seen = [], set()
    for i, h in enumerate(header):
        name = re.sub(r"[^a-z0-9]+", "_", (clean_text(h) or f"col_{i + 1}").lower()).strip("_") or f"col_{i + 1}"
        if name[0].isdigit():
            name = "c_" + name
        while name in seen:
            name += "_2"
        seen.add(name)
        cols.append((i, name))
    used = [(i, n) for i, n in cols if any(clean_text(cell(r, i)) for r in rows[1:])]
    if not used:
        return None
    tname = re.sub(r"[^a-z0-9]+", "_", ws.title.lower()).strip("_")
    t = Table(f"sheet_{tname}", f"Varaq '{ws.title.strip()}' (avtomatik yuklangan)", [(n, "TEXT", "") for _, n in used])
    for r in rows[1:]:
        t.rows.append({n: (to_date(cell(r, i)) if isinstance(cell(r, i), dt.datetime) else clean_text(cell(r, i))) for i, n in used})
    return t


_MAKE_FIX = {"Freighliner": "Freightliner", "Freightliner Freightliner": "Freightliner"}


def _fix_truck_row(row: dict) -> None:
    for k in TRUCK_COLUMNS:
        row.setdefault(k[0], None)
    if row["color"]:
        row["color"] = row["color"].title().replace("Grey", "Gray").replace("Golden", "Gold")
    if row["make"]:
        row["make"] = _MAKE_FIX.get(row["make"], row["make"])
    if row["model"]:
        row["model"] = re.sub(r"(?i)^lt\s*625$", "LT 625", row["model"])
    # miles_until_pm ni o'zimiz hisoblaymiz: jadvaldagi formula bo'sh probegda noto'g'ri qiymat beradi
    nxt, cur = row["next_pm_due_miles"], row["current_mileage"]
    row["miles_until_pm"] = nxt - cur if (nxt and cur) else None


TRUCK_SHEETS = {
    "TRUCK INFO": parse_standard_trucks,
    "CAC Trucks": parse_standard_trucks,
    "DEMIR Trucks": parse_standard_trucks,
    "New trucks": parse_new_trucks,
    "Rental trucks": parse_rental,
    "Inactive Trucks": parse_inactive,
}
OTHER_SHEETS = {
    "Sheet103": parse_recent_issues,
    "Insurance VIN": parse_telematics,
    "PTI": parse_pti,
    "Sheet45": parse_contacts,
    "2022": parse_deductions,
    "Sheet35": parse_days_off,
}
# 2022 yilgi eski shablon / bo'sh varaqlar — yuklanmaydi
SKIP_SHEETS = {"Fleet Live", "Odometer", "TR Detention", "Sheet for use", "Sheet87", "Copy of Sheet87", "Sheet90"}

TRUCKS_DESCRIPTION = (
    "Barcha yuk mashinalari (6 ta varaq birlashtirilgan). 'TRUCK INFO' = hozirgi asosiy aktiv park. "
    "'CAC Trucks' va 'DEMIR Trucks' = alohida ro'yxatlar, 'New trucks' = yangi mashinalar (probeg 2026-04-16 holatiga), "
    "'Rental trucks' = ijaradagi mashinalar, 'Inactive Trucks' = faol bo'lmagan/qaytarilgan mashinalar arxivi. "
    "Bitta mashina bir necha varaqda bo'lishi mumkin."
)


# --------------------------------------------------------------------------- #
# Yuklash
# --------------------------------------------------------------------------- #

@dataclass
class LoadedData:
    conn: sqlite3.Connection
    tables: list[Table]
    source: str
    loaded_at: dt.datetime


def load_workbook_tables(path: str | Path) -> list[Table]:
    with warnings.catch_warnings():  # "Cell O100 is marked as a date..." — bitta buzuq katak, e'tiborsiz
        warnings.simplefilter("ignore", UserWarning)
        wb = openpyxl.load_workbook(path, data_only=True, read_only=False)
    trucks = Table("all_trucks", TRUCKS_DESCRIPTION, TRUCK_COLUMNS)
    others: list[Table] = []
    for ws in wb.worksheets:
        title = ws.title.strip()
        if title in SKIP_SHEETS:
            continue
        if title in TRUCK_SHEETS:
            trucks.rows.extend(TRUCK_SHEETS[title](ws, title))
        elif title in OTHER_SHEETS:
            others.append(OTHER_SHEETS[title](ws))
        else:
            t = parse_generic(ws)
            if t:
                others.append(t)
    for row in trucks.rows:
        _fix_truck_row(row)
    return [trucks] + [t for t in others if t.rows]


def build_sqlite(tables: list[Table]) -> sqlite3.Connection:
    conn = sqlite3.connect(":memory:", check_same_thread=False)
    for t in tables:
        cols_sql = ", ".join(f'"{n}" {typ}' for n, typ, _ in t.columns)
        conn.execute(f'CREATE TABLE "{t.name}" ({cols_sql})')
        names = [n for n, _, _ in t.columns]
        conn.executemany(
            f'INSERT INTO "{t.name}" VALUES ({", ".join("?" * len(names))})',
            [[row.get(n) for n in names] for row in t.rows],
        )
    conn.execute("CREATE INDEX idx_all_trucks_unit ON all_trucks(unit_number)")
    for name, sql in TRUCK_VIEWS.items():
        conn.execute(f"CREATE VIEW {name} AS {sql}")
    conn.commit()
    conn.execute("PRAGMA query_only = ON")
    return conn


# LLM ko'radigan view'lar. `trucks` = faqat aktiv park, shuning uchun oddiy COUNT(*) to'g'ri son beradi.
TRUCK_VIEWS = {
    "trucks": "SELECT driver, coworker, location_note, unit_number, vehicle_group, owner, color, make, model, year, vin, "
              "plate_number, prepass, ifta, cab_card_state, cab_card_expiration, inspection_state, inspection_expiration, "
              "latest_pm_miles, next_pm_due_miles, miles_until_pm, current_mileage FROM all_trucks WHERE sheet = 'TRUCK INFO'",
    "rental_trucks": "SELECT driver, coworker, location_note, unit_number, vehicle_group, make, model, year, vin "
                     "FROM all_trucks WHERE sheet = 'Rental trucks'",
    "other_truck_lists": "SELECT sheet AS list_name, driver, coworker, location_note, unit_number, vehicle_group, owner, make, "
                         "model, year, vin, plate_number, cab_card_expiration, inspection_expiration, current_mileage, "
                         "miles_until_pm FROM all_trucks WHERE sheet IN ('CAC Trucks', 'DEMIR Trucks', 'New trucks')",
    "inactive_trucks": "SELECT driver AS last_driver, location_note, unit_number, COALESCE(owner, vehicle_group) AS provider, "
                       "make, model, year, vin, plate_number, cab_card_expiration, inspection_expiration "
                       "FROM all_trucks WHERE sheet = 'Inactive Trucks'",
}


def load(path: str | Path) -> LoadedData:
    tables = load_workbook_tables(path)
    return LoadedData(build_sqlite(tables), tables, str(path), dt.datetime.now())


def _top(conn, sql: str, n: int = 12) -> str:
    rows = conn.execute(sql).fetchall()
    s = ", ".join(f"'{v}'({c})" for v, c in rows[:n] if v is not None)
    return s + (", ..." if len(rows) > n else "")


def _count(conn, table: str) -> int:
    return conn.execute(f'SELECT COUNT(*) FROM "{table}"').fetchone()[0]


def schema_text(data: LoadedData) -> str:
    """LLM uchun ixcham schema (har chaqiruvda yuboriladi, shuning uchun token tejaladi)."""
    c = data.conn
    names = {t.name for t in data.tables}
    out = [
        f"trucks — ACTIVE FLEET, {_count(c, 'trucks')} trucks (the main 'TRUCK INFO' sheet). Use it for every question about "
        "the fleet/our trucks/counts/PM/inspections/drivers unless the user names another list.",
        "  unit_number TEXT (e.g. '417', '9693 JCE', '8578 RH')",
        "  driver TEXT, coworker TEXT (team driver). driver IS NULL = no driver; then location_note says where the truck stands",
        "  location_note TEXT: " + _top(c, "SELECT location_note, COUNT(*) FROM trucks WHERE location_note IS NOT NULL GROUP BY 1 ORDER BY 2 DESC"),
        "  vehicle_group TEXT: " + _top(c, "SELECT vehicle_group, COUNT(*) FROM trucks GROUP BY 1 ORDER BY 2 DESC", 10),
        "  owner TEXT, color TEXT, year INTEGER",
        "  make TEXT: " + _top(c, "SELECT make, COUNT(*) FROM trucks GROUP BY 1 ORDER BY 2 DESC")
        + "; model TEXT: " + _top(c, "SELECT model, COUNT(*) FROM trucks GROUP BY 1 ORDER BY 2 DESC", 6),
        "  vin, plate_number, prepass, ifta TEXT; cab_card_state, inspection_state TEXT (US state code)",
        "  cab_card_expiration, inspection_expiration TEXT 'YYYY-MM-DD' (cab card = registration)",
        "  current_mileage, latest_pm_miles, next_pm_due_miles INTEGER (miles)",
        "  miles_until_pm INTEGER = next_pm_due_miles - current_mileage; < 0 means PM overdue by that many miles",
        f"rental_trucks — {_count(c, 'rental_trucks')} rented trucks, NOT part of the active fleet count: "
        "driver, coworker, location_note, unit_number, vehicle_group, make, model, year, vin",
        "other_truck_lists — side lists " + _top(c, "SELECT list_name, COUNT(*) FROM other_truck_lists GROUP BY 1")
        + "; mostly the same trucks as `trucks` — never add them to fleet counts. Columns: list_name, driver, coworker, "
        "location_note, unit_number, vehicle_group, owner, make, model, year, vin, plate_number, cab_card_expiration, "
        "inspection_expiration, current_mileage, miles_until_pm",
        f"inactive_trucks — ARCHIVE of {_count(c, 'inactive_trucks')} old/returned trucks; use only when asked about "
        "inactive/returned/old trucks: last_driver, location_note, unit_number, provider, make, model, year, vin, "
        "plate_number, cab_card_expiration, inspection_expiration",
    ]
    if "recent_issues" in names:
        out.append(f"recent_issues — {_count(c, 'recent_issues')} rows, latest repair/issue per truck: unit_number, driver, "
                   "location_note, make, model, year, vin, mileage, recent_issue")
    if "gps_tracking" in names:
        last = c.execute("SELECT MAX(last_connected) FROM gps_tracking").fetchone()[0]
        out.append(f"gps_tracking — {_count(c, 'gps_tracking')} rows, Samsara GPS export (latest signal {last}): unit_number, "
                   "location, last_connected, last_trip ('On trip' = driving now), engine_status ('on','off','idle'), "
                   "fuel_level_pct REAL, current_driver, vin")
    if "pti_checks" in names:
        out.append("pti_checks — daily PTI (pre-trip inspection) requests: drivers, note, check_date 'YYYY-MM-DD', status: "
                   + _top(c, "SELECT status, COUNT(*) FROM pti_checks GROUP BY 1"))
    if "contacts" in names:
        out.append("contacts — office staff: name, email, phone")
    if "staff_days_off" in names:
        out.append("staff_days_off — office staff weekly day off: name, day_off")
    if "deductions_2022" in names:
        out.append("deductions_2022 — old 2022 broker deductions: reason, date, load_number, amount_deducted, load_amount")
    known = {"all_trucks", "recent_issues", "gps_tracking", "pti_checks", "contacts", "staff_days_off", "deductions_2022"}
    for t in data.tables:  # jadvalga keyin qo'shilgan varaqlar
        if t.name not in known:
            out.append(f"{t.name} — {t.description}: " + ", ".join(n for n, _, _ in t.columns))
    return "\n".join(out)
